import re
import pdfplumber


import os

import docx

def extract_text_from_file(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in [".txt", ".md", ".json"]:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return clean_text(f.read())
    elif ext in [".docx", ".doc"]:
        try:
            doc = docx.Document(path)
            full_text = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_txt:
                        full_text.append(row_txt)
            return clean_text("\n".join(full_text))
        except Exception as e:
            print(f"docx parsing fallback for {path}: {e}")
            try:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    return clean_text(f.read())
            except Exception:
                return ""
    else:
        text = ""
        try:
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception:
            try:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
            except Exception:
                text = ""
        return clean_text(text)


def extract_text_from_pdf(path: str) -> str:
    return extract_text_from_file(path)


def clean_text(text: str) -> str:
    text = text.replace("\uf0b7", " ")
    text = text.replace("●", " ")
    text = re.sub(r'[^\S\n]+', ' ', text)  # collapse spaces but KEEP newlines
    text = re.sub(r'\n{3,}', '\n\n', text)  # max 2 consecutive newlines
    return text.strip()